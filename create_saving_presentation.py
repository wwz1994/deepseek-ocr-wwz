#!/usr/bin/env python3
"""
Create a comprehensive presentation about saving money.
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor


def add_title_slide(prs, title, subtitle):
    """Add a title slide to the presentation."""
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    
    title_shape = slide.shapes.title
    subtitle_shape = slide.placeholders[1]
    
    title_shape.text = title
    subtitle_shape.text = subtitle
    
    # Format title
    title_frame = title_shape.text_frame
    title_frame.paragraphs[0].font.size = Pt(44)
    title_frame.paragraphs[0].font.bold = True
    title_frame.paragraphs[0].font.color.rgb = RGBColor(0, 51, 102)
    
    return slide


def add_content_slide(prs, title, content_items):
    """Add a content slide with bullet points."""
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    
    title_shape = slide.shapes.title
    body_shape = slide.placeholders[1]
    
    title_shape.text = title
    
    # Format title
    title_frame = title_shape.text_frame
    title_frame.paragraphs[0].font.size = Pt(32)
    title_frame.paragraphs[0].font.bold = True
    title_frame.paragraphs[0].font.color.rgb = RGBColor(0, 51, 102)
    
    # Add content
    text_frame = body_shape.text_frame
    text_frame.clear()
    
    for i, item in enumerate(content_items):
        if i == 0:
            p = text_frame.paragraphs[0]
        else:
            p = text_frame.add_paragraph()
        
        p.text = item
        p.level = 0
        p.font.size = Pt(18)
        p.space_after = Pt(12)
    
    return slide


def create_saving_presentation():
    """Create a comprehensive presentation about saving."""
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)
    
    # Slide 1: Title
    add_title_slide(
        prs,
        "The Art of Saving",
        "Building Financial Security for Your Future"
    )
    
    # Slide 2: Why Saving Matters
    add_content_slide(
        prs,
        "Why Saving Matters",
        [
            "Financial Security: Protection against unexpected expenses",
            "Peace of Mind: Reduced stress about money matters",
            "Future Goals: Ability to achieve dreams and aspirations",
            "Independence: Freedom from debt and financial constraints",
            "Opportunities: Flexibility to seize life's chances"
        ]
    )
    
    # Slide 3: Types of Savings
    add_content_slide(
        prs,
        "Types of Savings",
        [
            "Emergency Fund: 3-6 months of living expenses",
            "Short-term Savings: Goals within 1-3 years (vacation, car)",
            "Medium-term Savings: Goals within 3-10 years (home down payment)",
            "Long-term Savings: Retirement and distant future needs",
            "Specialized Savings: Education, healthcare, major purchases"
        ]
    )
    
    # Slide 4: The 50/30/20 Rule
    add_content_slide(
        prs,
        "The 50/30/20 Budgeting Rule",
        [
            "50% - Needs: Essential expenses (housing, food, utilities)",
            "30% - Wants: Discretionary spending (entertainment, dining)",
            "20% - Savings: Emergency fund, retirement, debt repayment",
            "Adjust percentages based on your personal situation",
            "Track spending to ensure you're meeting your targets"
        ]
    )
    
    # Slide 5: Effective Saving Strategies
    add_content_slide(
        prs,
        "Effective Saving Strategies",
        [
            "Pay Yourself First: Automate savings before spending",
            "Set Clear Goals: Define specific, measurable targets",
            "Track Expenses: Know where your money goes",
            "Reduce Unnecessary Costs: Cut subscriptions and impulse buys",
            "Increase Income: Side hustles and skill development",
            "Use High-Yield Accounts: Maximize interest earnings"
        ]
    )
    
    # Slide 6: Building an Emergency Fund
    add_content_slide(
        prs,
        "Building an Emergency Fund",
        [
            "Start Small: Begin with $500-$1,000 as initial goal",
            "Calculate Needs: 3-6 months of essential expenses",
            "Keep It Accessible: Use savings account, not investments",
            "Separate Account: Don't mix with regular spending money",
            "Replenish After Use: Rebuild fund after emergencies",
            "Review Annually: Adjust for life changes and inflation"
        ]
    )
    
    # Slide 7: Common Saving Mistakes
    add_content_slide(
        prs,
        "Common Saving Mistakes to Avoid",
        [
            "No Clear Goals: Saving without purpose leads to failure",
            "Lifestyle Inflation: Spending more as income increases",
            "Ignoring Small Expenses: Daily coffee adds up over time",
            "Not Automating: Relying on willpower instead of systems",
            "Keeping All in Checking: Missing out on interest earnings",
            "Dipping Into Savings: Using savings for non-emergencies"
        ]
    )
    
    # Slide 8: Savings vs. Investment
    add_content_slide(
        prs,
        "Savings vs. Investment",
        [
            "Savings: Low risk, liquid, lower returns (0.5-5% annually)",
            "Investments: Higher risk, potential higher returns (7-10% avg)",
            "Emergency Fund: Always keep in savings",
            "Short-term Goals: Use savings accounts",
            "Long-term Goals: Consider investments after emergency fund",
            "Diversification: Balance between safety and growth"
        ]
    )
    
    # Slide 9: Tools and Apps
    add_content_slide(
        prs,
        "Tools to Help You Save",
        [
            "Budgeting Apps: Mint, YNAB, Personal Capital",
            "Automatic Savings: Digit, Qapital, Acorns",
            "High-Yield Accounts: Online banks offering better rates",
            "Expense Trackers: Monitor spending patterns",
            "Goal Calculators: Determine how much to save monthly",
            "Financial Advisors: Professional guidance when needed"
        ]
    )
    
    # Slide 10: Getting Started Today
    add_content_slide(
        prs,
        "Getting Started Today",
        [
            "Step 1: Calculate your monthly income and expenses",
            "Step 2: Set one specific savings goal",
            "Step 3: Open a dedicated savings account",
            "Step 4: Automate a small weekly or monthly transfer",
            "Step 5: Track progress and celebrate milestones",
            "Remember: Small consistent actions lead to big results!"
        ]
    )
    
    # Slide 11: Conclusion
    add_content_slide(
        prs,
        "Key Takeaways",
        [
            "Saving is a habit that can be learned and improved",
            "Start small and be consistent",
            "Automate your savings to remove temptation",
            "Have clear goals to stay motivated",
            "Review and adjust your strategy regularly",
            "Your future self will thank you for starting today!"
        ]
    )
    
    # Save the presentation
    output_file = "Saving_Presentation.pptx"
    prs.save(output_file)
    print(f"✓ Presentation created successfully: {output_file}")
    print(f"✓ Total slides: {len(prs.slides)}")
    return output_file


if __name__ == "__main__":
    create_saving_presentation()
